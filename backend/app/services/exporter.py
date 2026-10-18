from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas import ExpenseOut

RUPIAH = "#,##0"
TANGGAL = "dd mmm yyyy"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="059669") 
BOLD = Font(bold=True)


def export_filename(start: date | None, end: date | None) -> str:
    if start and end:
        return f"pengeluaran_{start}_{end}.xlsx"
    if start:
        return f"pengeluaran_sejak_{start}.xlsx"
    if end:
        return f"pengeluaran_sampai_{end}.xlsx"
    return "pengeluaran.xlsx"


def _text(ws, row: int, col: int, value: str):
    cell = ws.cell(row=row, column=col, value=ILLEGAL_CHARACTERS_RE.sub("", value))
    cell.data_type = "s"
    return cell


def _number(ws, row: int, col: int, value, bold: bool = False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.number_format = RUPIAH
    if bold:
        cell.font = BOLD
    return cell


def _date(ws, row: int, col: int, value: date):
    cell = ws.cell(row=row, column=col, value=value)
    cell.number_format = TANGGAL
    return cell


def _header(ws, titles: list[str], widths: list[int]) -> None:
    for col, (title, width) in enumerate(zip(titles, widths), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"


def build_workbook(expenses: list[ExpenseOut]) -> bytes:
    ordered = sorted(expenses, key=lambda e: (e.expense_date, e.id))

    wb = Workbook()
    summary = wb.active
    summary.title = "Transaksi"
    _header(
        summary,
        ["Tanggal", "Barang", "Jumlah barang (Rp)", "Penyesuaian (Rp)", "Total dibayar (Rp)"],
        [14, 50, 20, 18, 20],
    )

    detail = wb.create_sheet("Rincian barang")
    _header(
        detail,
        ["Tanggal", "Nama barang", "Jumlah", "Harga satuan (Rp)", "Subtotal (Rp)"],
        [14, 50, 10, 20, 20],
    )

    sum_items = sum_adjustment = sum_total = Decimal(0)
    detail_row = 2

    for row, e in enumerate(ordered, start=2):
        _date(summary, row, 1, e.expense_date)
        _text(summary, row, 2, ", ".join(i.item_name for i in e.items))
        _number(summary, row, 3, e.items_total)
        _number(summary, row, 4, e.adjustment)
        _number(summary, row, 5, e.total)
        sum_items += e.items_total
        sum_adjustment += e.adjustment
        sum_total += e.total

        for item in e.items:
            _date(detail, detail_row, 1, e.expense_date)
            _text(detail, detail_row, 2, item.item_name)
            detail.cell(row=detail_row, column=3, value=item.quantity)
            _number(detail, detail_row, 4, item.price)
            _number(detail, detail_row, 5, item.subtotal)
            detail_row += 1
            
    last = len(ordered) + 2
    _text(summary, last, 2, "Total").font = BOLD
    _number(summary, last, 3, sum_items, bold=True)
    _number(summary, last, 4, sum_adjustment, bold=True)
    _number(summary, last, 5, sum_total, bold=True)
    
    if ordered:
        summary.auto_filter.ref = f"A1:E{last - 1}"
        detail.auto_filter.ref = f"A1:E{detail_row - 1}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()